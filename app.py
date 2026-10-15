import os
from pathlib import Path
import shutil
import tempfile
import json
from flask import Flask, request, jsonify
import yake
from pypdf import PdfReader
import docx
from pptx import Presentation

app = Flask(__name__)

@app.after_request
def after_request(response):
    response.headers.add('Access-Control-Allow-Origin', '*')
    response.headers.add('Access-Control-Allow-Headers', 'Content-Type,Authorization')
    response.headers.add('Access-Control-Allow-Methods', 'GET,PUT,POST,DELETE,OPTIONS')
    return response

def extract_text_from_file(file_path: str) -> str:
    ext = Path(file_path).suffix.lower()
    text = ""
    if ext == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
    elif ext == ".pdf":
        reader = PdfReader(file_path)
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    elif ext in [".docx", ".doc"]:
        doc = docx.Document(file_path)
        paragraphs = [para.text for para in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.append(cell.text)
        text = "\n".join(paragraphs)
    elif ext == ".pptx":
        prs = Presentation(file_path)
        slide_texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        slide_texts.append(paragraph.text)
        text = "\n".join(slide_texts)
    else:
        raise ValueError(f"Unsupported file format: {ext}")
    return text

@app.route('/extract-keywords', methods=['POST'])
def extract_keywords():
    if 'file' not in request.files:
        return jsonify({"error": "No file provided"}), 400
    file = request.files['file']
    supported_extensions = ['.txt', '.pdf', '.docx', '.doc', '.pptx']
    ext = Path(file.filename).suffix.lower()
    if ext not in supported_extensions:
        return jsonify({"detail": f"Unsupported file format: {ext}"}), 400

    unwanted_words_json = request.form.get('unwanted_words_json', '[]')
    try:
        extra_unwanted = json.loads(unwanted_words_json)
        if isinstance(extra_unwanted, list):
            unwanted_words = [
                'chapter', 'summary', 'overview', 'slep', 'contents',
                'introduction', 'conclusion', 'declaration', 'acknowledgement'
            ] + [str(w).lower() for w in extra_unwanted]
        else:
            unwanted_words = [
                'chapter', 'summary', 'overview', 'slep', 'contents',
                'introduction', 'conclusion', 'declaration', 'acknowledgement'
            ]
    except Exception:
        unwanted_words = [
            'chapter', 'summary', 'overview', 'slep', 'contents',
            'introduction', 'conclusion', 'declaration', 'acknowledgement'
        ]

    with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as temp_file:
        file.save(temp_file)
        temp_file_path = temp_file.name

    try:
        document_text = extract_text_from_file(temp_file_path)
        if not document_text.strip():
            return jsonify({"filename": file.filename, "keywords": [], "message": "Document is empty or text could not be extracted."})

        custom_kw_extractor = yake.KeywordExtractor(
            lan="en",
            n=3,
            dedupLim=0.9,
            top=20
        )
        raw_keywords = custom_kw_extractor.extract_keywords(document_text)
        filtered_keywords = []
        for kw, score in raw_keywords:
            if not any(unwanted in kw.lower() for unwanted in unwanted_words):
                filtered_keywords.append({"keyword": kw, "score": round(float(score), 4)})
        filtered_keywords = filtered_keywords[:10]
        return jsonify({"filename": file.filename, "keywords": filtered_keywords})
    except Exception as e:
        return jsonify({"detail": str(e)}), 500
    finally:
        if os.path.exists(temp_file_path):
            os.remove(temp_file_path)

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=8000)
